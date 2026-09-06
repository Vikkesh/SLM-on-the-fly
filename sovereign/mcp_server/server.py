"""HTTP MCP server exposing the tools TrueForge does not ship with: OCR a scan, and write real
Word / PDF / Excel documents from one shared structure (title, metadata, sections with optional
tables and bullets, sign-off).

TrueForge's MCP client speaks Streamable HTTP / SSE only (no stdio), so this runs as a small web
service on localhost and is registered under Settings -> Connectors by URL.
"""

from __future__ import annotations

import os
import re
import time
from pathlib import Path

from mcp.server.fastmcp import FastMCP
from mcp.types import ToolAnnotations
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path(os.environ.get("OUTPUT_DIR", ROOT / "output"))
UPLOAD_DIR = Path(os.environ.get("UPLOAD_DIR", ROOT / "uploads"))
HOST = os.environ.get("MCP_HOST", "127.0.0.1")
PORT = int(os.environ.get("MCP_PORT", "9000"))

mcp = FastMCP("sovereign-tools", host=HOST, port=PORT)

READ = ToolAnnotations(readOnlyHint=True, destructiveHint=False, openWorldHint=False)
WRITE = ToolAnnotations(readOnlyHint=False, destructiveHint=False, idempotentHint=False, openWorldHint=False)


# --- shared document structure -------------------------------------------------------------------


class Table(BaseModel):
    headers: list[str] = Field(description="Column headings, e.g. ['Parameter', 'Reading', 'Limit', 'Status']")
    rows: list[list[str]] = Field(description="Rows of cell text, same length as headers")


class Section(BaseModel):
    heading: str = Field(description="Section title, e.g. 'Summary of findings'")
    body: str = Field(default="", description="Paragraph text. Blank lines separate paragraphs.")
    bullets: list[str] = Field(default_factory=list, description="Optional bullet points after the body")
    table: Table | None = Field(default=None, description="Optional table after the bullets")


DOC_DESCRIPTION = (
    " Provide a title, an optional metadata table (e.g. Asset ID, Inspector, Date, Risk rating), ordered "
    "sections (paragraphs, bullets, an optional table each) and optional sign-off roles. Returns the saved path."
)


# --- extract_from_scan ---------------------------------------------------------------------------


@mcp.tool(annotations=READ, description="Run Tesseract OCR on an uploaded scan (PNG/JPG under the uploads folder) and return the exact text.")
def extract_from_scan(image_path: str) -> str:
    path = _inside(UPLOAD_DIR, image_path)
    try:
        import pytesseract
        from PIL import Image
    except ImportError as e:  # pragma: no cover
        return f"OCR unavailable: {e}"
    try:
        text = pytesseract.image_to_string(Image.open(path))
    except pytesseract.TesseractNotFoundError:
        return "OCR unavailable: the tesseract binary is not installed on this machine (apt install tesseract-ocr)."
    return text.strip() or "(no text recognised)"


# --- generate_docx -------------------------------------------------------------------------------


@mcp.tool(annotations=WRITE, description="Create a formal Word (.docx) document." + DOC_DESCRIPTION)
def generate_docx(
    title: str,
    sections: list[Section],
    metadata: dict[str, str] | None = None,
    signoff: list[str] | None = None,
    filename: str | None = None,
) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if metadata:
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for key, value in metadata.items():
            cells = t.add_row().cells
            cells[0].text, cells[1].text = str(key), str(value)
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
        doc.add_paragraph()

    for s in sections:
        doc.add_heading(s.heading, level=1)
        for para in (p.strip() for p in s.body.split("\n\n")):
            if para:
                doc.add_paragraph(para)
        for b in s.bullets:
            doc.add_paragraph(b, style="List Bullet")
        if s.table and s.table.headers:
            t = doc.add_table(rows=1, cols=len(s.table.headers))
            t.style = "Light Grid Accent 1"
            for cell, head in zip(t.rows[0].cells, s.table.headers):
                cell.text = str(head)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in s.table.rows:
                cells = t.add_row().cells
                for cell, val in zip(cells, list(row) + [""] * (len(s.table.headers) - len(row))):
                    cell.text = str(val)
                    if str(val).strip().upper() in ("OVER", "EXCEEDED", "FAIL", "HIGH"):
                        for run in cell.paragraphs[0].runs:
                            run.font.color.rgb = RGBColor(0xB0, 0x1E, 0x1E)
                            run.bold = True
            doc.add_paragraph()

    if signoff:
        doc.add_heading("Sign-off", level=1)
        t = doc.add_table(rows=1, cols=len(signoff))
        t.style = "Table Grid"
        for cell, role in zip(t.rows[0].cells, signoff):
            cell.text = f"{role}\n\nName: ____________\nSignature: ________\nDate: ____________"

    path = _target(filename or title, ".docx")
    doc.save(path)
    return f"Saved: {path} ({path.stat().st_size} bytes)"


# --- generate_pdf --------------------------------------------------------------------------------


@mcp.tool(annotations=WRITE, description="Create a formal PDF document." + DOC_DESCRIPTION)
def generate_pdf(
    title: str,
    sections: list[Section],
    metadata: dict[str, str] | None = None,
    signoff: list[str] | None = None,
    filename: str | None = None,
) -> str:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    h0 = ParagraphStyle("h0", parent=styles["Title"], alignment=TA_LEFT, fontSize=18, leading=22, spaceAfter=10)
    h1 = ParagraphStyle("h1", parent=styles["Heading2"], fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=15)
    small = ParagraphStyle("small", parent=body, fontSize=9.5, leading=13)
    grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9AA3AD")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8ECEF")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])
    P = lambda text, st=body: Paragraph(_pdf_escape(str(text)), st)  # noqa: E731

    story = [P(title, h0)]
    if metadata:
        rows = [[P(k, small), P(v, small)] for k, v in metadata.items()]
        t = Table(rows, colWidths=[50 * mm, 110 * mm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9AA3AD")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8ECEF")),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story += [t, Spacer(1, 6)]

    for s in sections:
        story.append(P(s.heading, h1))
        for para in (p.strip() for p in s.body.split("\n\n")):
            if para:
                story.append(P(para))
        if s.bullets:
            story.append(ListFlowable([ListItem(P(b), leftIndent=12) for b in s.bullets], bulletType="bullet", start="•", leftIndent=14))
        if s.table and s.table.headers:
            width = len(s.table.headers)
            data = [[P(h, small) for h in s.table.headers]]
            style = TableStyle(grid.getCommands())
            for r_i, row in enumerate(s.table.rows, start=1):
                cells = list(row) + [""] * (width - len(row))
                data.append([P(c, small) for c in cells[:width]])
                for c_i, c in enumerate(cells[:width]):
                    if str(c).strip().upper() in ("OVER", "EXCEEDED", "FAIL", "HIGH"):
                        style.add("TEXTCOLOR", (c_i, r_i), (c_i, r_i), colors.HexColor("#B01E1E"))
            t = Table(data, repeatRows=1, colWidths=[(160 * mm) / width] * width)
            t.setStyle(style)
            story += [Spacer(1, 4), t]

    if signoff:
        story.append(P("Sign-off", h1))
        cells = [Paragraph(f"<b>{_pdf_escape(role)}</b><br/><br/>Name: ____________<br/>Signature: ________<br/>Date: ____________", small) for role in signoff]
        t = Table([cells], colWidths=[(160 * mm) / len(signoff)] * len(signoff))
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9AA3AD")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [Spacer(1, 6), t]

    path = _target(filename or title, ".pdf")
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm, topMargin=22 * mm,
                      bottomMargin=22 * mm, title=title).build(story)
    return f"Saved: {path} ({path.stat().st_size} bytes)"


# --- generate_xlsx -------------------------------------------------------------------------------


class Sheet(BaseModel):
    name: str = Field(description="Sheet name (max 31 chars)")
    headers: list[str]
    rows: list[list[str | float | int | None]] = Field(default_factory=list)


@mcp.tool(annotations=WRITE, description="Create an Excel (.xlsx) workbook. Either pass headers+rows for one sheet, or `sheets` for several. Returns the saved path.")
def generate_xlsx(
    headers: list[str] | None = None,
    rows: list[list[str | float | int | None]] | None = None,
    sheet_name: str = "Sheet1",
    sheets: list[Sheet] | None = None,
    filename: str | None = None,
) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    todo = sheets or [Sheet(name=sheet_name, headers=headers or [], rows=rows or [])]
    if not todo or not todo[0].headers:
        raise ValueError("Provide headers and rows, or sheets")
    head_fill = PatternFill("solid", fgColor="E8ECEF")
    for sh in todo:
        ws = wb.create_sheet(sh.name[:31])
        ws.append(sh.headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for r in sh.rows:
            ws.append([_num(v) for v in r])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(ws.columns, start=1):
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(10, width + 2), 60)
    path = _target(filename or todo[0].name, ".xlsx")
    wb.save(path)
    return f"Saved: {path} ({path.stat().st_size} bytes)"


# --- list_outputs --------------------------------------------------------------------------------


@mcp.tool(annotations=READ, description="List documents generated so far (newest first) with their paths.")
def list_outputs(limit: int = 20) -> str:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted((p for p in OUTPUT_DIR.iterdir() if p.is_file() and not p.name.startswith(".")), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        return "(no documents generated yet)"
    return "\n".join(f"{p} ({p.stat().st_size} bytes)" for p in files[:limit])


# --- helpers -------------------------------------------------------------------------------------


def _num(v):
    if isinstance(v, str):
        try:
            f = float(v.replace(",", ""))
            return int(f) if f.is_integer() and "." not in v else f
        except ValueError:
            return v
    return v


def _pdf_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def _target(name: str, ext: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", Path(name).stem).strip("-")[:60] or "document"
    return OUTPUT_DIR / f"{stem}-{time.strftime('%Y%m%d-%H%M%S')}{ext}"


def _inside(folder: Path, requested: str) -> Path:
    """Only files under the uploads folder may be read; a bare filename is resolved there."""
    p = Path(requested)
    candidate = (p if p.is_absolute() else folder / p).resolve()
    if folder.resolve() not in candidate.parents and candidate != folder.resolve():
        raise ValueError(f"Path must be inside {folder}")
    if not candidate.is_file():
        raise FileNotFoundError(f"No such upload: {requested}")
    return candidate


def main() -> None:
    mcp.run(transport="streamable-http")


if __name__ == "__main__":
    main()
