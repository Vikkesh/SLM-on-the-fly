"""Turn any upload into what the models can actually consume: text, or PNG images.

The models on Laptop A only ever see text and images. TrueForge would send PDFs in a shape Ollama
rejects and drop everything else into the sandbox as a filename stub, so we convert up front.
"""

from __future__ import annotations

import csv
import io
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

from . import config

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
TEXT_EXT = {".txt", ".md", ".json", ".log", ".yaml", ".yml"}


class UnsupportedFile(Exception):
    pass


@dataclass
class TextPart:
    label: str
    text: str


@dataclass
class ImagePart:
    label: str
    png: bytes
    saved_path: Path = field(default=None)  # type: ignore[assignment]


Part = TextPart | ImagePart


def normalize(filename: str, data: bytes) -> list[Part]:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXT:
        return [_image(filename, data)]
    if ext == ".pdf":
        return _pdf(filename, data)
    if ext in {".xlsx", ".xlsm"}:
        return [_xlsx(filename, data)]
    if ext == ".csv":
        return [_csv(filename, data)]
    if ext == ".docx":
        return [_docx(filename, data)]
    if ext in TEXT_EXT:
        return [TextPart(filename, _cap(data.decode("utf-8", "replace")))]
    raise UnsupportedFile(f"Unsupported file type: {ext or filename}")


# --- images -------------------------------------------------------------------------------------


def _image(filename: str, data: bytes) -> ImagePart:
    from PIL import Image

    png = _to_png(Image.open(io.BytesIO(data)))
    return ImagePart(filename, png, _save(Path(filename).stem, png))


def _to_png(img) -> bytes:
    """Normalize mode and cap the long edge: vision-model cost scales with pixel count."""
    from PIL import Image

    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    if max(img.size) > config.MAX_IMAGE_EDGE:
        img.thumbnail((config.MAX_IMAGE_EDGE, config.MAX_IMAGE_EDGE), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def _save(stem: str, png: bytes) -> Path:
    """Persist the PNG so the MCP `extract_from_scan` tool can OCR it by path."""
    config.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", stem)[:40] or "upload"
    path = config.UPLOAD_DIR / f"{safe}-{int(time.time() * 1000)}.png"
    path.write_bytes(png)
    return path


# --- pdf ----------------------------------------------------------------------------------------


def _pdf(filename: str, data: bytes) -> list[Part]:
    """Digital PDFs become text; scanned ones become one PNG per page (capped)."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    total = sum(len(p.strip()) for p in pages)
    if pages and total / len(pages) >= config.SCANNED_PDF_CHARS_PER_PAGE:
        text = "\n\n".join(f"[page {i + 1}]\n{p.strip()}" for i, p in enumerate(pages) if p.strip())
        return [TextPart(filename, _cap(text))]
    return _rasterize(filename, data, len(pages))


def _rasterize(filename: str, data: bytes, page_count: int) -> list[Part]:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(data)
    parts: list[Part] = []
    limit = min(page_count, config.MAX_PDF_PAGES)
    for i in range(limit):
        png = _to_png(doc[i].render(scale=config.PDF_RENDER_SCALE).to_pil())
        label = f"{filename} (page {i + 1}/{page_count})"
        parts.append(ImagePart(label, png, _save(f"{Path(filename).stem}-p{i + 1}", png)))
    if page_count > limit:
        parts.append(TextPart(filename, f"[note] {page_count} pages in this PDF; only the first {limit} were sent."))
    return parts


# --- tables -------------------------------------------------------------------------------------


def _xlsx(filename: str, data: bytes) -> TextPart:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks = []
    for ws in wb.worksheets:
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v != "" for v in r)]
        chunks.append(f"### Sheet: {ws.title}\n" + _table(rows))
    return TextPart(filename, _cap("\n\n".join(chunks)))


def _csv(filename: str, data: bytes) -> TextPart:
    text = data.decode("utf-8", "replace")
    rows = [row for row in csv.reader(io.StringIO(text)) if any(v.strip() for v in row)]
    return TextPart(filename, _cap(_table(rows)))


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\n", " ").replace("|", "\\|")


def _table(rows: list[list[str]]) -> str:
    if not rows:
        return "(empty)"
    header, body = rows[0], rows[1:]
    shown = body[: config.MAX_TABLE_ROWS]
    width = max(len(r) for r in rows)
    pad = lambda r: [str(v) for v in r] + [""] * (width - len(r))  # noqa: E731
    lines = ["| " + " | ".join(pad(header)) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(pad(r)) + " |" for r in shown]
    shape = f"{len(body)} data rows x {width} columns"
    if len(body) > len(shown):
        shape += f", showing first {len(shown)} - ask for a specific range for more"
    return f"({shape})\n" + "\n".join(lines)


# --- docx ---------------------------------------------------------------------------------------


def _docx(filename: str, data: bytes) -> TextPart:
    from docx import Document

    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
        out.append(_table(rows))
    return TextPart(filename, _cap("\n\n".join(out)))


def _cap(text: str) -> str:
    if len(text) <= config.MAX_TEXT_CHARS:
        return text
    return text[: config.MAX_TEXT_CHARS] + f"\n\n[truncated: {len(text) - config.MAX_TEXT_CHARS} more characters]"
